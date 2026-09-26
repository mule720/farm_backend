"""Donor results report tests (PDF + Excel)."""
import json
import os
import tempfile
from datetime import date, timedelta

from django.contrib.auth.models import AnonymousUser
from django.test import RequestFactory, TestCase, override_settings

from apps.accounts.models import Organization, Profile
from apps.enterprises.models import Enterprise, EnterpriseBatch
from apps.production.models import ProductionRecord
from apps.partners.models import Programme, ProgrammeEnrollment, ProgrammeSupport, IndicatorReading
from apps.partners import reports
from config.schema import schema

TMP = tempfile.mkdtemp(prefix='agrinuxes_reports_')


def _gql(user, query, variables=None):
    request = RequestFactory().post('/graphql/', data=json.dumps({'query': query}), content_type='application/json')
    request.user = user if user is not None else AnonymousUser()
    return schema.execute(query, variable_values=variables or {}, context_value=request)


@override_settings(MEDIA_ROOT=TMP, MEDIA_URL='/media/')
class ProgrammeReportTest(TestCase):
    def setUp(self):
        self.partner = Organization.objects.create(name='FAO Zambia', slug='fao-r', org_type='donor')
        self.pm = Profile.objects.create_user(email='pm@fao.r', full_name='PM', password='pw12345678', phone='0955700001', organization=self.partner, role='partner_manager')
        self.observer = Profile.objects.create_user(email='ob@fao.r', full_name='Observer', password='pw12345678', phone='0955700002', organization=self.partner, role='partner_observer')
        other = Organization.objects.create(name='WFP', slug='wfp-r', org_type='ngo')
        self.other = Profile.objects.create_user(email='o@wfp.r', full_name='O', password='pw12345678', phone='0955700003', organization=other, role='partner_manager')
        self.prog = Programme.objects.create(organization=self.partner, name='Hand-in-Hand Zambia', code='GCP/ZAM/101', funder='FAO', status='active',
                                             start_date=date.today() - timedelta(days=60), budget=100000, currency='USD', target_farms=4,
                                             target_provinces=['Lusaka'], description='Raise maize productivity.',
                                             indicators=[{'key': 'yield', 'label': 'Maize yield', 'unit': 't/ha', 'target': 3.5}, {'key': 'trained', 'label': 'Farmers trained', 'unit': 'farmers', 'target': 100}],
                                             created_by=self.pm)
        for i, (name, dist) in enumerate([('Farm A', 'Chongwe'), ('Farm B', 'Kafue')]):
            org = Organization.objects.create(name=name, slug=f'rf-{i}', org_type='farm', province='Lusaka', district=dist, data_sharing_consent=True)
            d = Profile.objects.create_user(email=f'd{i}@rf.r', full_name='D', password='pw12345678', phone=f'097770000{i}', organization=org, role='director')
            ent = Enterprise.objects.create(organization=org, name='Maize', category='crop', production_type='maize', created_by=d)
            b = EnterpriseBatch.objects.create(organization=org, enterprise=ent, name='B', start_date=date.today(), status='active', created_by=d)
            ProductionRecord.objects.create(organization=org, enterprise=ent, batch=b, record_date=date.today() - timedelta(days=5), record_type='harvest', data={'quantity': 300}, recorded_by=d)
            ProductionRecord.objects.create(organization=org, enterprise=ent, batch=b, record_date=date.today() - timedelta(days=90), record_type='harvest', data={'quantity': 999}, recorded_by=d)  # before programme
            ProgrammeEnrollment.objects.create(programme=self.prog, farm=org, cohort='2026A', enrolled_by=self.pm)
            ProgrammeSupport.objects.create(programme=self.prog, farm=org, support_type='input_voucher', description='Inputs', value=1500, currency='USD', delivered_on=date.today() - timedelta(days=10), reference=f'V-{i}', created_by=self.pm)
        IndicatorReading.objects.create(programme=self.prog, indicator_key='yield', period=date.today() - timedelta(days=30), value=2.1, recorded_by=self.pm)
        IndicatorReading.objects.create(programme=self.prog, indicator_key='yield', period=date.today(), value=2.8, notes='crop cut', recorded_by=self.pm)

    def test_collect_data(self):
        d = reports.collect_report_data(self.prog)
        self.assertEqual(d['results'].farms_enrolled, 2)
        self.assertEqual(len(d['districts']), 2)
        self.assertAlmostEqual(sum(x['harvest'] for x in d['districts']), 600.0)  # the pre-programme harvest is excluded
        self.assertEqual(d['support_by_type'][0]['events'], 2)
        self.assertAlmostEqual(d['support_by_type'][0]['value'], 3000.0)
        y = next(i for i in d['indicators'] if i['key'] == 'yield')
        self.assertAlmostEqual(y['latest'], 2.8)
        self.assertAlmostEqual(y['progress_pct'], 80.0)
        self.assertEqual(len(y['history']), 2)
        t = next(i for i in d['indicators'] if i['key'] == 'trained')
        self.assertIsNone(t['latest'])

    def test_period_filter(self):
        d = reports.collect_report_data(self.prog, period_start=date.today() - timedelta(days=7), period_end=date.today())
        self.assertAlmostEqual(sum(x['harvest'] for x in d['districts']), 600.0)
        self.assertEqual(d['ledger'], [])  # support was 10 days ago

    def _check(self, url, ext, magic):
        self.assertTrue(url.startswith('/media/reports/') and url.endswith(ext), url)
        path = os.path.join(TMP, 'reports', os.path.basename(url))
        self.assertTrue(os.path.exists(path))
        with open(path, 'rb') as f:
            self.assertEqual(f.read(len(magic)), magic)
        self.assertGreater(os.path.getsize(path), 3000)
        return path

    def test_pdf_and_xlsx_render(self):
        self._check(reports.build_programme_report_pdf(self.prog), '.pdf', b'%PDF-')
        path = self._check(reports.build_programme_report_xlsx(self.prog), '.xlsx', b'PK')
        import openpyxl
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames, ['Summary', 'Indicators', 'Indicator readings', 'Districts', 'Enrolled farms', 'Support by type', 'Support ledger'])
        self.assertEqual(wb['Enrolled farms'].max_row, 3)
        self.assertEqual(wb['Support ledger'].max_row, 3)
        self.assertEqual(wb['Indicators']['G2'].value, 80.0)

    def test_empty_programme_renders(self):
        empty = Programme.objects.create(organization=self.partner, name='Empty', start_date=date.today(), created_by=self.pm)
        self._check(reports.build_programme_report_pdf(empty), '.pdf', b'%PDF-')
        self._check(reports.build_programme_report_xlsx(empty), '.xlsx', b'PK')

    def test_mutation_access_and_formats(self):
        m = 'mutation($id: ID!, $f: String, $s: Date, $e: Date) { generateProgrammeReport(programmeId: $id, format: $f, periodStart: $s, periodEnd: $e) { url } }'
        r = _gql(self.pm, m, {'id': str(self.prog.id)})
        self.assertIsNone(r.errors, r.errors)
        self._check(r.data['generateProgrammeReport']['url'], '.pdf', b'%PDF-')
        r = _gql(self.observer, m, {'id': str(self.prog.id), 'f': 'xlsx'})
        self.assertIsNone(r.errors, r.errors)
        self._check(r.data['generateProgrammeReport']['url'], '.xlsx', b'PK')
        self.assertTrue(_gql(self.other, m, {'id': str(self.prog.id)}).errors)
        self.assertTrue(_gql(None, m, {'id': str(self.prog.id)}).errors)
        self.assertTrue(_gql(self.pm, m, {'id': str(self.prog.id), 'f': 'docx'}).errors)
        self.assertTrue(_gql(self.pm, m, {'id': str(self.prog.id), 's': '2026-06-01', 'e': '2026-01-01'}).errors)
